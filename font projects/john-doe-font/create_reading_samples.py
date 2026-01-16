"""
Create Word Document Reading Samples for John Doe Font DOE Testing
==================================================================

This script generates Word (.docx) documents for the Phase 1 DOE testing
of the John Doe personalized font optimization project.

Prerequisites:
- python-docx library installed (pip install python-docx)
- Font variants must be installed on the system:
  * ODx-Font1_Control
  * ODx-Font2_LetterSpace  
  * ODx-Font3_WordSpace
  * ODx-Font4_GlyphWeight
  * ODx-Font5_Combination

Output:
- 5 sample documents (A-E) in font/variants/font test doc files/
- Instructions cover page
- Results summary page
"""

from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

# Configuration
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "font", "variants", "font test doc files")
FONT_SIZE = Pt(16)
HEADING_SIZE = Pt(18)

# Sample mapping based on randomization key from JOHN_DOE_FONT_PROJECT.md
SAMPLES = {
    'A': {
        'font_name': 'ODx-Font3_WordSpace',
        'description': 'Word Spacing +40%',
        'file_name': 'Sample_A_WordSpace.docx'
    },
    'B': {
        'font_name': 'ODx-Font1_Control',
        'description': 'Baseline (Control)',
        'file_name': 'Sample_B_Baseline.docx'
    },
    'C': {
        'font_name': 'ODx-Font5_Combination',
        'description': 'Combination (Letter +25%, Word +40%, Heavy)',
        'file_name': 'Sample_C_Combination.docx'
    },
    'D': {
        'font_name': 'ODx-Font2_LetterSpace',
        'description': 'Letter Spacing +25%',
        'file_name': 'Sample_D_LetterSpace.docx'
    },
    'E': {
        'font_name': 'ODx-Font4_GlyphWeight',
        'description': 'Heavy Glyph Weight',
        'file_name': 'Sample_E_GlyphWeight.docx'
    }
}

# The reading passage (Max the dog story) - 3rd grade reading level
READING_PASSAGE = """Max was a small brown dog who lived on a farm. Every morning, he would wake up early and run outside to play. He loved to chase the chickens around the yard. The chickens would cluck and flap their wings, but Max never hurt them. He just wanted to have fun.

One sunny day, Max found a big red ball near the barn. He picked it up in his mouth and ran to show his best friend, a cat named Whiskers. Whiskers was sleeping in the hay, but she woke up when Max dropped the ball at her feet.

"Do you want to play?" Max asked with his tail wagging fast. Whiskers stretched and yawned. She was not sure about the ball, but she liked Max. So she gave the ball a little tap with her paw. It rolled away, and both friends chased after it together.

They played until the sun went down. Then they walked back to the farmhouse side by side. Max was happy. He had a good friend and a fun day. That night, they both slept very well."""


def set_paragraph_font(paragraph, font_name, font_size):
    """Set font for all runs in a paragraph"""
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = font_size


def create_sample_document(sample_letter, sample_info):
    """Create a single sample document"""
    doc = Document()
    
    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    
    font_name = sample_info['font_name']
    
    # Title
    title = doc.add_paragraph()
    title_run = title.add_run(f"SAMPLE {sample_letter}")
    title_run.font.name = 'Arial'
    title_run.font.size = HEADING_SIZE
    title_run.font.bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add baseline indicator for Sample B
    if sample_letter == 'B':
        subtitle = doc.add_paragraph()
        subtitle_run = subtitle.add_run("(Baseline Reference)")
        subtitle_run.font.name = 'Arial'
        subtitle_run.font.size = Pt(14)
        subtitle_run.font.bold = True
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()  # Spacer
    
    # Reading passage - split into paragraphs
    paragraphs = READING_PASSAGE.strip().split('\n\n')
    for para_text in paragraphs:
        para = doc.add_paragraph()
        run = para.add_run(para_text)
        run.font.name = font_name
        run.font.size = FONT_SIZE
        para.paragraph_format.space_after = Pt(12)
        para.paragraph_format.line_spacing = 1.6
    
    doc.add_paragraph()  # Spacer
    
    # Scoring section
    scoring_header = doc.add_paragraph()
    scoring_run = scoring_header.add_run("After reading, rate this sample:")
    scoring_run.font.name = 'Arial'
    scoring_run.font.size = Pt(12)
    scoring_run.font.bold = True
    
    if sample_letter == 'B':
        instruction = doc.add_paragraph()
        inst_run = instruction.add_run("This is the BASELINE (Score = 5). Compare other samples to this one.")
        inst_run.font.name = 'Arial'
        inst_run.font.size = Pt(11)
    else:
        instruction = doc.add_paragraph()
        inst_run = instruction.add_run("How easy was this to read compared to Sample B? (Circle one)")
        inst_run.font.name = 'Arial'
        inst_run.font.size = Pt(11)
    
    # Score scale
    scale = doc.add_paragraph()
    scale_run = scale.add_run("1     2     3     4     5     6     7     8     9     10")
    scale_run.font.name = 'Arial'
    scale_run.font.size = Pt(14)
    scale.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    labels = doc.add_paragraph()
    labels_run = labels.add_run("(Harder)                              (Baseline)                              (Easier)")
    labels_run.font.name = 'Arial'
    labels_run.font.size = Pt(10)
    labels.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()  # Spacer
    
    # Notes section
    notes_label = doc.add_paragraph()
    notes_run = notes_label.add_run("Notes (optional):")
    notes_run.font.name = 'Arial'
    notes_run.font.size = Pt(11)
    
    notes_line = doc.add_paragraph()
    line_run = notes_line.add_run("_" * 70)
    line_run.font.name = 'Arial'
    line_run.font.size = Pt(11)
    
    notes_line2 = doc.add_paragraph()
    line_run2 = notes_line2.add_run("_" * 70)
    line_run2.font.name = 'Arial'
    line_run2.font.size = Pt(11)
    
    return doc


def create_instructions_page():
    """Create the cover page with instructions"""
    doc = Document()
    
    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    
    # Title
    title = doc.add_paragraph()
    title_run = title.add_run("Font Readability Test")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Subtitle
    subtitle = doc.add_paragraph()
    sub_run = subtitle.add_run("Dyslexia Optimization Study - Phase 1")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(16)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Instructions header
    inst_header = doc.add_paragraph()
    inst_run = inst_header.add_run("Instructions for John")
    inst_run.font.name = 'Arial'
    inst_run.font.size = Pt(14)
    inst_run.font.bold = True
    
    # Instructions list
    instructions = [
        "Read each passage - There are 5 samples (A through E), each showing the same text in a different font style.",
        "Take your time - Read at a comfortable pace. This is not a speed test.",
        "Rate each sample - After reading, score how easy or difficult it was to read on a scale of 1-10.",
        "Use the rubric - The scoring guide is below. Sample B is the baseline (OpenDyslexic = 5).",
        "Be honest - There are no right or wrong answers. We want your genuine experience."
    ]
    
    for i, inst in enumerate(instructions, 1):
        para = doc.add_paragraph()
        run = para.add_run(f"{i}. {inst}")
        run.font.name = 'Arial'
        run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    # Scoring rubric header
    rubric_header = doc.add_paragraph()
    rubric_run = rubric_header.add_run("Scoring Scale (1-10)")
    rubric_run.font.name = 'Arial'
    rubric_run.font.size = Pt(14)
    rubric_run.font.bold = True
    
    # Scoring rubric table
    table = doc.add_table(rows=7, cols=2)
    table.style = 'Table Grid'
    
    rubric_data = [
        ("Score", "Meaning"),
        ("1-2", "Very difficult to read, causes significant strain or confusion"),
        ("3-4", "Harder to read than usual, requires extra effort"),
        ("5", "Baseline (OpenDyslexic) - Your current comfortable level"),
        ("6-7", "Noticeably easier than baseline, comfortable"),
        ("8-9", "Significantly easier, very comfortable reading"),
        ("10", "Ideal - effortless reading experience")
    ]
    
    for i, (score, meaning) in enumerate(rubric_data):
        row = table.rows[i]
        cell0 = row.cells[0]
        cell1 = row.cells[1]
        
        # Score column
        p0 = cell0.paragraphs[0]
        r0 = p0.add_run(score)
        r0.font.name = 'Arial'
        r0.font.size = Pt(11)
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Meaning column
        p1 = cell1.paragraphs[0]
        r1 = p1.add_run(meaning)
        r1.font.name = 'Arial'
        r1.font.size = Pt(11)
        
        # Header row styling
        if i == 0:
            r0.font.bold = True
            r1.font.bold = True
        
        # Highlight baseline row
        if score == "5":
            r0.font.bold = True
            r1.font.bold = True
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Date/Time fields
    date_para = doc.add_paragraph()
    date_run = date_para.add_run("Date: _____________     Time Started: _____________")
    date_run.font.name = 'Arial'
    date_run.font.size = Pt(12)
    
    return doc


def create_results_summary():
    """Create the results summary page"""
    doc = Document()
    
    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
    
    # Title
    title = doc.add_paragraph()
    title_run = title.add_run("Results Summary")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph()
    
    # Results table
    table = doc.add_table(rows=6, cols=3)
    table.style = 'Table Grid'
    
    headers = ["Sample", "Your Score (1-10)", "Notes"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        p = cell.paragraphs[0]
        r = p.add_run(header)
        r.font.name = 'Arial'
        r.font.size = Pt(11)
        r.font.bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    samples = [
        ("A", "", ""),
        ("B (Baseline)", "5", "OpenDyslexic standard"),
        ("C", "", ""),
        ("D", "", ""),
        ("E", "", "")
    ]
    
    for i, (sample, score, notes) in enumerate(samples, 1):
        row = table.rows[i]
        
        # Sample column
        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(sample)
        r0.font.name = 'Arial'
        r0.font.size = Pt(11)
        r0.font.bold = True
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Score column
        p1 = row.cells[1].paragraphs[0]
        r1 = p1.add_run(score)
        r1.font.name = 'Arial'
        r1.font.size = Pt(11)
        if score == "5":
            r1.font.bold = True
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Notes column
        p2 = row.cells[2].paragraphs[0]
        r2 = p2.add_run(notes)
        r2.font.name = 'Arial'
        r2.font.size = Pt(11)
    
    doc.add_paragraph()
    doc.add_paragraph()
    
    # Final Questions
    questions_header = doc.add_paragraph()
    q_run = questions_header.add_run("Final Questions")
    q_run.font.name = 'Arial'
    q_run.font.size = Pt(14)
    q_run.font.bold = True
    
    # Question 1
    q1 = doc.add_paragraph()
    q1_run = q1.add_run("1. Which sample was EASIEST to read? (Circle one)")
    q1_run.font.name = 'Arial'
    q1_run.font.size = Pt(12)
    q1_run.font.bold = True
    
    a1 = doc.add_paragraph()
    a1_run = a1.add_run("     A          B          C          D          E")
    a1_run.font.name = 'Arial'
    a1_run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    # Question 2
    q2 = doc.add_paragraph()
    q2_run = q2.add_run("2. Which sample was HARDEST to read? (Circle one)")
    q2_run.font.name = 'Arial'
    q2_run.font.size = Pt(12)
    q2_run.font.bold = True
    
    a2 = doc.add_paragraph()
    a2_run = a2.add_run("     A          B          C          D          E")
    a2_run.font.name = 'Arial'
    a2_run.font.size = Pt(12)
    
    doc.add_paragraph()
    
    # Question 3
    q3 = doc.add_paragraph()
    q3_run = q3.add_run("3. Any other comments about the reading experience?")
    q3_run.font.name = 'Arial'
    q3_run.font.size = Pt(12)
    q3_run.font.bold = True
    
    for _ in range(3):
        line = doc.add_paragraph()
        line_run = line.add_run("_" * 70)
        line_run.font.name = 'Arial'
        line_run.font.size = Pt(11)
    
    doc.add_paragraph()
    
    # Time completed
    time_para = doc.add_paragraph()
    time_run = time_para.add_run("Time Completed: _____________     Total Time: _____________ minutes")
    time_run.font.name = 'Arial'
    time_run.font.size = Pt(12)
    
    return doc


def main():
    """Main function to generate all documents"""
    
    # Ensure output directory exists
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    
    print("=" * 60)
    print("John Doe Font Project - Reading Sample Generator")
    print("=" * 60)
    print()
    
    # Create sample documents
    print("Creating sample documents...")
    for letter, info in SAMPLES.items():
        doc = create_sample_document(letter, info)
        filepath = os.path.join(OUTPUT_DIR, info['file_name'])
        doc.save(filepath)
        print(f"  ✓ Created: {info['file_name']} ({info['description']})")
    
    print()
    
    # Create instructions page
    print("Creating supporting documents...")
    instructions_doc = create_instructions_page()
    instructions_path = os.path.join(OUTPUT_DIR, "00_Instructions_Cover.docx")
    instructions_doc.save(instructions_path)
    print(f"  ✓ Created: 00_Instructions_Cover.docx")
    
    # Create results summary
    summary_doc = create_results_summary()
    summary_path = os.path.join(OUTPUT_DIR, "99_Results_Summary.docx")
    summary_doc.save(summary_path)
    print(f"  ✓ Created: 99_Results_Summary.docx")
    
    print()
    print("=" * 60)
    print("COMPLETE! All documents created successfully.")
    print("=" * 60)
    print()
    print(f"Output location: {OUTPUT_DIR}")
    print()
    print("IMPORTANT: Before printing, ensure these fonts are installed:")
    for letter, info in SAMPLES.items():
        print(f"  • {info['font_name']}")
    print()
    print("To install fonts: Right-click each .otf file → Install")
    print("Font files are in: font/variants/ and font/variants/weighted/")


if __name__ == "__main__":
    main()
